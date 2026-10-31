"""In-memory builders for real files of every supported Drive format (no network, no disk). Holds no tests."""
from __future__ import annotations

import io
import json
import zipfile
from typing import Dict, List, Optional


def make_text_pdf(pages: List[str]) -> bytes:
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas

    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=A4)
    for text in pages:
        y = 800
        for line in text.splitlines():
            c.drawString(50, y, line)
            y -= 16
        c.showPage()
    c.save()
    return buf.getvalue()


def make_png(size=(200, 80), color=(255, 255, 255)) -> bytes:
    from PIL import Image, ImageDraw

    img = Image.new("RGB", size, color)
    ImageDraw.Draw(img).text((10, 10), "Price 9,900,000 THB", fill=(0, 0, 0))
    out = io.BytesIO()
    img.save(out, format="PNG")
    return out.getvalue()


def make_scanned_pdf(pages: int = 1) -> bytes:
    """A PDF whose pages are only an embedded image (no text layer), like a scan."""
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.utils import ImageReader
    from reportlab.pdfgen import canvas

    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=A4)
    for _ in range(pages):
        c.drawImage(ImageReader(io.BytesIO(make_png((400, 160)))), 50, 600, width=400, height=160)
        c.showPage()
    c.save()
    return buf.getvalue()


def make_mixed_pdf() -> bytes:
    """Page 1 has text, page 2 is image-only."""
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.utils import ImageReader
    from reportlab.pdfgen import canvas

    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=A4)
    c.drawString(50, 800, "Property Name: Mixed Page Villa")
    c.showPage()
    c.drawImage(ImageReader(io.BytesIO(make_png((400, 160)))), 50, 600, width=400, height=160)
    c.showPage()
    c.save()
    return buf.getvalue()


def make_encrypted_pdf() -> bytes:
    import pypdf

    writer = pypdf.PdfWriter()
    reader = pypdf.PdfReader(io.BytesIO(make_text_pdf(["Secret Villa"])))
    for page in reader.pages:
        writer.add_page(page)
    writer.encrypt(user_password="pw", owner_password="owner")
    out = io.BytesIO()
    writer.write(out)
    return out.getvalue()


def make_docx(paragraphs: List[str], *, heading: Optional[str] = None, kv_table: Optional[List[List[str]]] = None,
              table: Optional[List[List[str]]] = None, header_text: Optional[str] = None,
              footer_text: Optional[str] = None) -> bytes:
    import docx

    doc = docx.Document()
    if header_text:
        doc.sections[0].header.paragraphs[0].text = header_text
    if footer_text:
        doc.sections[0].footer.paragraphs[0].text = footer_text
    if heading:
        doc.add_heading(heading, level=1)
    for p in paragraphs:
        doc.add_paragraph(p)
    for rows in (kv_table, table):
        if rows:
            t = doc.add_table(rows=len(rows), cols=len(rows[0]))
            for r, row in enumerate(rows):
                for c, val in enumerate(row):
                    t.cell(r, c).text = val
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def make_xlsx(sheets: Dict[str, List[List[object]]]) -> bytes:
    import openpyxl

    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for name, rows in sheets.items():
        ws = wb.create_sheet(name)
        for row in rows:
            ws.append(row)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def make_pptx(slides: List[Dict[str, object]]) -> bytes:
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    for spec in slides:
        slide = prs.slides.add_slide(prs.slide_layouts[5])  # title only
        slide.shapes.title.text = str(spec.get("title") or "")
        if spec.get("text"):
            box = slide.shapes.add_textbox(Inches(1), Inches(2), Inches(6), Inches(2))
            box.text_frame.text = str(spec["text"])
        rows = spec.get("table")
        if rows:
            shape = slide.shapes.add_table(len(rows), len(rows[0]), Inches(0.5), Inches(3), Inches(8), Inches(2))
            for r, row in enumerate(rows):
                for c, val in enumerate(row):
                    shape.table.cell(r, c).text = str(val)
        if spec.get("notes"):
            slide.notes_slide.notes_text_frame.text = str(spec["notes"])
        if spec.get("picture"):
            slide.shapes.add_picture(io.BytesIO(make_png()), Inches(1), Inches(5))
    out = io.BytesIO()
    prs.save(out)
    return out.getvalue()


def make_zip(members: Dict[str, bytes], *, compression=zipfile.ZIP_DEFLATED) -> bytes:
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", compression=compression) as zf:
        for name, data in members.items():
            zf.writestr(name, data)
    return out.getvalue()


def make_rtf(text: str) -> bytes:
    body = text.replace("\n", "\\par\n")
    return ("{\\rtf1\\ansi\\deff0{\\fonttbl{\\f0 Arial;}}{\\*\\generator Test;}\\f0 " + body + "}").encode("latin-1")


def ooxml_encrypted_stub() -> bytes:
    """Encrypted Office files are OLE compound files, not ZIPs."""
    return b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 504


def listing_rows() -> List[List[object]]:
    return [
        ["Property ID", "Property Name", "Type", "Status", "Province", "District", "Price (THB)", "Bedrooms",
         "Bathrooms", "Area (sqm)", "Furnishing", "Amenities", "Image URLs"],
        ["kamala-sky-villa", "Kamala Sky Villa", "Villa", "Available", "Phuket", "Kamala", 45000000, 4, 4, 420,
         "Fully furnished", "Private pool, Gym, Sea view terrace", "https://cdn.example.com/k1.jpg"],
        ["lamai-garden-condo", "Lamai Garden Condo", "Condo", "Under construction", "Surat Thani", "Lamai",
         "8,500,000", "1-2", 1, "45-68", None, "Pool; Parking", None],
        [None, None, None, None, None, None, None, None, None, None, None, None, None],
    ]


def json_listings() -> bytes:
    return json.dumps({"properties": [
        {"name": "Bang Tao Beach House", "propertyType": "House", "status": "Ready to move in",
         "address": {"addressRegion": "Phuket", "addressLocality": "Bang Tao"},
         "price": "29,000,000 THB", "bedrooms": 3, "bathrooms": 3, "floorSize": "310 sqm"},
        {"name": "Chaweng Studio", "type": "Condominium", "status": "Pre-sale", "province": "Surat Thani",
         "price": 3500000, "currency": "THB", "bedrooms": "studio"},
    ]}).encode("utf-8")


def xml_listings() -> bytes:
    return b"""<?xml version="1.0"?>
<feed><listings>
  <property id="rawai-1"><name>Rawai Hill Villa</name><type>Villa</type><status>Available</status>
    <province>Phuket</province><price currency="THB">18500000</price><bedrooms>3</bedrooms></property>
  <property id="rawai-2"><name>Rawai Hill Villa B</name><type>Villa</type><status>Planned</status>
    <province>Phuket</province><price currency="USD">600000</price><bedrooms>4</bedrooms></property>
</listings></feed>"""


# ---- in-memory Firestore double (same behaviour as the URL-ingest tests' fake) ---------------------------------------
import copy as _copy
import uuid as _uuid


class _Snap:
    def __init__(self, doc_id, data):
        self.id = doc_id
        self.exists = data is not None
        self._data = None if data is None else _copy.deepcopy(data)

    def to_dict(self):
        return None if self._data is None else _copy.deepcopy(self._data)


class _Doc:
    def __init__(self, store, collection, doc_id):
        self.store, self.collection_name, self.id = store, collection, doc_id

    def get(self):
        return _Snap(self.id, self.store.get(self.collection_name, self.id))

    def set(self, data, merge=False):
        self.store.set(self.collection_name, self.id, data, merge=merge)

    def update(self, data):
        self.store.update(self.collection_name, self.id, data)

    def delete(self):
        self.store.delete(self.collection_name, self.id)


class _Query:
    def __init__(self, collection):
        self.collection = collection

    def limit(self, _n):
        return self

    def where(self, *a, **k):
        return self

    def stream(self):
        return self.collection.stream()


class _Collection:
    def __init__(self, store, name):
        self.store, self.name = store, name

    def document(self, doc_id=None):
        return _Doc(self.store, self.name, str(doc_id or _uuid.uuid4().hex))

    def stream(self):
        for doc_id, data in self.store.items(self.name):
            yield _Snap(doc_id, data)

    def limit(self, _n):
        return _Query(self)

    def where(self, *a, **k):
        return _Query(self)


class FakeFirestore:
    def __init__(self):
        self.docs = {}

    def collection(self, name):
        return _Collection(self, name)

    def get(self, collection, doc_id):
        return self.docs.get((collection, doc_id))

    def items(self, collection):
        return [(d, v) for (c, d), v in list(self.docs.items()) if c == collection]

    def set(self, collection, doc_id, data, merge=False):
        key = (collection, doc_id)
        payload = _copy.deepcopy(data)
        current = self.docs.get(key)
        if merge and isinstance(current, dict):
            merged = _copy.deepcopy(current)
            merged.update(payload)
            self.docs[key] = merged
        else:
            self.docs[key] = payload

    def update(self, collection, doc_id, data):
        current = self.docs.get((collection, doc_id))
        if not isinstance(current, dict):
            raise RuntimeError(doc_id)
        current = _copy.deepcopy(current)
        current.update(_copy.deepcopy(data))
        self.docs[(collection, doc_id)] = current

    def delete(self, collection, doc_id):
        self.docs.pop((collection, doc_id), None)

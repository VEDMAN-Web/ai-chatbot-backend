"""Generated report files open and carry the right content: PDF, Excel, Word and CSV (incl. Polish and Thai text)."""
import csv
import io
import re

import openpyxl
import pypdf
import docx
import pytest

from app.services import report_service as rs

LONG_NAME = "The Extraordinarily Long Residence Name Beyond Twenty Eight Characters"
ROWS = [
    {"property_id": "PHK-001", "property_name": "Layan Palm Villa", "location": "Layan, Phuket", "price_display": "from 8.5 million THB",
     "bedrooms": "2-3", "bathrooms": 2, "status": "Pre-sale", "developer_name": "Test Dev"},
    {"property_id": "SMU-002", "property_name": "Zażółć Gęślą Jaźń ทดสอบ", "location": "Maenam, Samui", "price_display": "12 000 000 THB",
     "bedrooms": "3", "bathrooms": 3, "status": "Ready", "developer_name": ""},
    {"property_id": "LNG-003", "property_name": LONG_NAME, "location": "A & B <Bay>", "price_display": "from 1 000 000 000 THB per unit",
     "bedrooms": "4", "bathrooms": 4, "status": "Under construction", "developer_name": ""},
]


def _flat(text: str) -> str:
    return re.sub(r"\s+", "", text)


def _pdf_text(data: bytes) -> str:
    assert data[:5] == b"%PDF-"
    return "".join(page.extract_text() or "" for page in pypdf.PdfReader(io.BytesIO(data)).pages)


def test_property_pdf_keeps_polish_thai_and_long_values():
    text = _flat(_pdf_text(rs._build_pdf_report(ROWS, "Inventory")))
    assert "Zażółć" in text and "Gęślą" in text and "Jaźń" in text
    assert "ทดสอบ" in text
    assert _flat(LONG_NAME) in text                     # not cut at 28 characters
    assert _flat("from 1 000 000 000 THB per unit") in text
    assert "A&B<Bay>" in text                           # markup characters are shown, not parsed


def test_property_pdf_with_no_rows_is_a_valid_file():
    assert "Total Properties: 0" in _pdf_text(rs._build_pdf_report([], "Empty"))


def test_lead_pdf_keeps_polish_thai_and_markup_characters():
    leads = [{"contact_name": "Zażółć & Co <VIP>", "contact_email": "a@b.com", "contact_phone": "+48 1", "preferred_location": "ภูเก็ต",
              "budget": "10M", "notes": "Chce <b>kupić</b> willę", "lead_score": 80, "status": "new", "buying_signals": ["Asked 'price'", "ทดสอบ"]}]
    text = _flat(_pdf_text(rs._build_lead_pdf_report(leads, "Leads")))
    assert "Zażółć&Co<VIP>" in text and "ภูเก็ต" in text and "<b>kupić</b>" in text and "ทดสอบ" in text


def test_excel_report_rows_match_the_data():
    ws = openpyxl.load_workbook(io.BytesIO(rs._build_excel_report(ROWS, "Inventory"))).active
    rows = [r for r in ws.iter_rows(min_row=3, values_only=True) if any(r)]
    assert [r[0] for r in rows] == ["PHK-001", "SMU-002", "LNG-003"]
    assert rows[1][1] == ROWS[1]["property_name"] and rows[0][3] == "from 8.5 million THB"


def test_word_report_contains_every_row():
    document = docx.Document(io.BytesIO(rs._build_docx_report(ROWS, "Inventory")))
    cells = [c.text for table in document.tables for row in table.rows for c in row.cells]
    for row in ROWS:
        assert row["property_id"] in cells
    assert ROWS[1]["property_name"] in "".join(cells)


def test_csv_report_round_trips():
    lines = list(csv.reader(io.StringIO(rs._build_csv_report(ROWS).decode("utf-8-sig"))))
    assert len(lines) == 4 and ROWS[1]["property_name"] in lines[2]


@pytest.mark.parametrize("builder", ["_build_excel_report", "_build_docx_report", "_build_pdf_report"])
def test_empty_and_large_datasets_produce_valid_files(builder):
    build = getattr(rs, builder)
    assert len(build([], "Empty")) > 500
    many = [dict(ROWS[0], property_id=f"BIG-{i:04d}") for i in range(300)]
    assert len(build(many, "Many")) > 5000


def test_missing_font_files_fall_back_without_failing(monkeypatch):
    monkeypatch.setattr(rs, "_PDF_FONTS", None)
    monkeypatch.setattr(rs, "_PDF_FONT_DIR", "/definitely/not/here")
    try:
        assert rs._register_pdf_fonts() == ("Helvetica", "Helvetica-Bold", False)
        assert _pdf_text(rs._build_pdf_report(ROWS[:1], "Fallback")).count("PHK-001") == 1
    finally:
        monkeypatch.setattr(rs, "_PDF_FONTS", None)  # let later tests register the real fonts again

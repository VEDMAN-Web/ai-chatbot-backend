"""
Format detection and per-format text/structure extraction for Google Drive bulk ingestion.

Every byte handled here is UNTRUSTED. Nothing is executed: no macros, no scripts, no embedded objects. Parsing runs in
memory (no temporary files are written), behind size, row, page, member and time limits.

Each extractor returns an ExtractionResult holding:
  - segments: text blocks with a source reference (page / sheet+row / slide / section / archive member)
  - records:  structured rows (spreadsheet rows, CSV rows, JSON objects, XML elements, document tables) with a reference
  - images:   metadata of embedded images (name, reference) — the image bytes are never stored
  - status:   ok | partial | needs_ocr | empty | error | unsupported | encrypted | too_large | timeout
"""
from __future__ import annotations

import csv
import io
import json
import logging
import os
import re
import time
import zipfile
from dataclasses import dataclass, field
from html.parser import HTMLParser
from typing import Any, Callable, Dict, List, Optional, Tuple

from app.core.config import settings

logger = logging.getLogger(__name__)

# --------------------------------------------------------------------------------------------------------------------
# Statuses
# --------------------------------------------------------------------------------------------------------------------
ST_OK = "ok"
ST_PARTIAL = "partial"
ST_NEEDS_OCR = "needs_ocr"
ST_EMPTY = "empty"
ST_ERROR = "error"
ST_UNSUPPORTED = "unsupported"
ST_ENCRYPTED = "encrypted"
ST_TOO_LARGE = "too_large"
ST_TIMEOUT = "timeout"
USABLE_STATUSES = (ST_OK, ST_PARTIAL)

# --------------------------------------------------------------------------------------------------------------------
# Formats and MIME types
# --------------------------------------------------------------------------------------------------------------------
GOOGLE_FOLDER = "application/vnd.google-apps.folder"
GOOGLE_DOC = "application/vnd.google-apps.document"
GOOGLE_SHEET = "application/vnd.google-apps.spreadsheet"
GOOGLE_SLIDES = "application/vnd.google-apps.presentation"
GOOGLE_SHORTCUT = "application/vnd.google-apps.shortcut"

MIME_DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
MIME_XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
MIME_PPTX = "application/vnd.openxmlformats-officedocument.presentationml.presentation"

# Google Workspace files have no bytes of their own; they are exported to an Office format and parsed with the same
# extractor, which keeps tables, headings, every worksheet and slide boundaries (CSV export would keep one sheet only).
GOOGLE_EXPORTS: Dict[str, Tuple[str, str]] = {
    GOOGLE_DOC: ("google_doc", MIME_DOCX),
    GOOGLE_SHEET: ("google_sheet", MIME_XLSX),
    GOOGLE_SLIDES: ("google_slides", MIME_PPTX),
}

# format key -> human label (used in counts by format)
FORMAT_LABELS: Dict[str, str] = {
    "pdf": "PDF",
    "docx": "Word (.docx)",
    "doc": "Word 97-2003 (.doc)",
    "google_doc": "Google Docs",
    "xlsx": "Excel (.xlsx)",
    "xls": "Excel 97-2003 (.xls)",
    "csv": "CSV",
    "google_sheet": "Google Sheets",
    "pptx": "PowerPoint (.pptx)",
    "ppt": "PowerPoint 97-2003 (.ppt)",
    "google_slides": "Google Slides",
    "txt": "Text (.txt)",
    "md": "Markdown (.md)",
    "rtf": "Rich Text (.rtf)",
    "html": "HTML",
    "json": "JSON",
    "xml": "XML",
    "image": "Image",
    "zip": "ZIP archive",
    "unsupported": "Unsupported",
}

_EXT_FORMAT: Dict[str, str] = {
    ".pdf": "pdf",
    ".docx": "docx", ".doc": "doc", ".docm": "docm", ".dotx": "docx",
    ".xlsx": "xlsx", ".xlsm": "xlsm", ".xls": "xls",
    ".csv": "csv", ".tsv": "csv",
    ".pptx": "pptx", ".pptm": "pptm", ".ppt": "ppt",
    ".txt": "txt", ".text": "txt", ".md": "md", ".markdown": "md", ".rtf": "rtf",
    ".html": "html", ".htm": "html",
    ".json": "json", ".xml": "xml",
    ".jpg": "image", ".jpeg": "image", ".png": "image", ".webp": "image", ".tif": "image", ".tiff": "image",
    ".gif": "image", ".bmp": "image",
    ".zip": "zip",
}

_MIME_FORMAT: Dict[str, str] = {
    "application/pdf": "pdf",
    MIME_DOCX: "docx",
    "application/msword": "doc",
    "application/vnd.ms-word.document.macroenabled.12": "docm",
    MIME_XLSX: "xlsx",
    "application/vnd.ms-excel": "xls",
    "application/vnd.ms-excel.sheet.macroenabled.12": "xlsm",
    "text/csv": "csv",
    "text/tab-separated-values": "csv",
    MIME_PPTX: "pptx",
    "application/vnd.ms-powerpoint": "ppt",
    "application/vnd.ms-powerpoint.presentation.macroenabled.12": "pptm",
    "text/plain": "txt",
    "text/markdown": "md",
    "text/x-markdown": "md",
    "application/rtf": "rtf",
    "text/rtf": "rtf",
    "text/html": "html",
    "application/xhtml+xml": "html",
    "application/json": "json",
    "application/xml": "xml",
    "text/xml": "xml",
    "image/jpeg": "image",
    "image/png": "image",
    "image/webp": "image",
    "image/tiff": "image",
    "image/gif": "image",
    "image/bmp": "image",
    "application/zip": "zip",
    "application/x-zip-compressed": "zip",
}

# Formats that are recognised but cannot be processed reliably. Reported as unsupported with this reason.
UNSUPPORTED_REASONS: Dict[str, str] = {
    "doc": "Legacy Word 97-2003 (.doc) files cannot be parsed reliably. Save it as .docx (or open it as a Google Doc) and re-run.",
    "ppt": "Legacy PowerPoint 97-2003 (.ppt) files cannot be parsed reliably. Save it as .pptx (or open it in Google Slides) and re-run.",
    "docm": "Macro-enabled Word files (.docm) are not processed. Macros are never run; save it as .docx and re-run.",
    "xlsm": "Macro-enabled Excel files (.xlsm) are not processed. Macros are never run; save it as .xlsx and re-run.",
    "pptm": "Macro-enabled PowerPoint files (.pptm) are not processed. Macros are never run; save it as .pptx and re-run.",
}

_GOOGLE_UNSUPPORTED = {
    "application/vnd.google-apps.form": "Google Forms have no exportable property content.",
    "application/vnd.google-apps.drawing": "Google Drawings are not processed.",
    "application/vnd.google-apps.site": "Google Sites cannot be exported through the Drive API.",
    "application/vnd.google-apps.map": "Google My Maps cannot be exported through the Drive API.",
    "application/vnd.google-apps.script": "Apps Script files are code and are never processed.",
    "application/vnd.google-apps.jam": "Jamboard files are not processed.",
    GOOGLE_SHORTCUT: "Drive shortcuts are not followed; add the target file to the folder instead.",
}

_EXECUTABLE_EXT = {".exe", ".dll", ".so", ".bin", ".bat", ".sh", ".cmd", ".com", ".msi", ".ps1", ".vbs", ".js", ".jar",
                   ".scr", ".app", ".apk", ".dmg", ".py", ".pyc"}

TEXT_FORMATS = {"txt", "md", "csv", "json", "xml", "html"}
OOXML_FORMATS = {"docx", "xlsx", "pptx", "google_doc", "google_sheet", "google_slides"}


def _cfg(name: str, default: Any) -> Any:
    value = getattr(settings, name, None)
    return default if value is None or value == "" else value


def max_file_bytes() -> int:
    return int(_cfg("DRIVE_INGEST_MAX_FILE_BYTES", 25 * 1024 * 1024))


def safe_filename(name: Optional[str], limit: int = 180) -> str:
    """Display/storage-safe file name: no path parts, no control characters, bounded length."""
    base = (name or "").replace("\\", "/").split("/")[-1]
    base = re.sub(r"[\x00-\x1f\x7f]", "", base).strip().strip(".") or "unnamed"
    return base[:limit]


def _ext(name: Optional[str]) -> str:
    return os.path.splitext((name or "").lower())[1]


@dataclass
class FormatInfo:
    format: str                 # key from FORMAT_LABELS (or docm/xlsm/pptm)
    category: str               # label used in counts by format
    supported: bool
    reason: Optional[str] = None
    export_mime: Optional[str] = None   # Google Workspace export target


def classify(name: Optional[str], mime_type: Optional[str]) -> FormatInfo:
    """Format from the Drive MIME type, then the extension. Bytes are verified later with verify_signature()."""
    mime = (mime_type or "").lower().split(";")[0].strip()
    ext = _ext(name)
    if mime in GOOGLE_EXPORTS:
        key, export = GOOGLE_EXPORTS[mime]
        return FormatInfo(key, FORMAT_LABELS[key], True, export_mime=export)
    if mime in _GOOGLE_UNSUPPORTED:
        return FormatInfo("unsupported", "Unsupported", False, _GOOGLE_UNSUPPORTED[mime])
    if mime.startswith("application/vnd.google-apps."):
        return FormatInfo("unsupported", "Unsupported", False, f"Google Workspace type '{mime}' cannot be exported for reading.")
    if ext in _EXECUTABLE_EXT or mime in ("application/x-msdownload", "application/x-executable", "application/x-sh",
                                          "application/x-msdos-program", "application/javascript", "text/javascript"):
        return FormatInfo("unsupported", "Unsupported", False, "Executable or script files are never processed.")
    key = _MIME_FORMAT.get(mime)
    # Drive often reports a generic type (octet-stream, text/plain for .md/.csv). The extension refines it.
    if key in (None, "txt") or mime in ("application/octet-stream", "binary/octet-stream"):
        key = _EXT_FORMAT.get(ext, key)
    if key is None:
        return FormatInfo("unsupported", "Unsupported", False,
                          f"No parser for this file type ({mime or 'unknown MIME type'}, extension '{ext or 'none'}').")
    if key in UNSUPPORTED_REASONS:
        label = FORMAT_LABELS.get(key, key)
        return FormatInfo(key, label if key in FORMAT_LABELS else "Unsupported", False, UNSUPPORTED_REASONS[key])
    return FormatInfo(key, FORMAT_LABELS[key], True)


# --------------------------------------------------------------------------------------------------------------------
# Signature checks
# --------------------------------------------------------------------------------------------------------------------
_OLE = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"


def sniff(data: bytes) -> str:
    """What the bytes really are: pdf, ooxml_docx/xlsx/pptx, ooxml_encrypted, ole, zip, png, jpeg, webp, tiff, gif, bmp,
    rtf, executable, text, binary or empty."""
    if not data:
        return "empty"
    head = data[:16]
    if b"%PDF-" in data[:1024]:
        return "pdf"
    if head.startswith(b"MZ") or head.startswith(b"\x7fELF") or head[:4] in (b"\xfe\xed\xfa\xce", b"\xcf\xfa\xed\xfe",
                                                                            b"\xca\xfe\xba\xbe"):
        return "executable"
    if head.startswith(b"#!"):
        return "executable"
    if head.startswith(_OLE):
        return "ole"
    if head.startswith(b"PK\x03\x04") or head.startswith(b"PK\x05\x06"):
        try:
            with zipfile.ZipFile(io.BytesIO(data)) as zf:
                names = set(zf.namelist())
        except Exception:
            return "zip"
        if "[Content_Types].xml" in names:
            if any(n.startswith("word/") for n in names):
                return "ooxml_docx"
            if any(n.startswith("xl/") for n in names):
                return "ooxml_xlsx"
            if any(n.startswith("ppt/") for n in names):
                return "ooxml_pptx"
        return "zip"
    if head.startswith(b"\x89PNG\r\n\x1a\n"):
        return "png"
    if head.startswith(b"\xff\xd8\xff"):
        return "jpeg"
    if head[:4] == b"RIFF" and data[8:12] == b"WEBP":
        return "webp"
    if head[:4] in (b"II*\x00", b"MM\x00*"):
        return "tiff"
    if head[:6] in (b"GIF87a", b"GIF89a"):
        return "gif"
    if head[:2] == b"BM":
        return "bmp"
    if head.lstrip().startswith(b"{\\rtf"):
        return "rtf"
    if _looks_like_text(data):
        return "text"
    return "binary"


def _looks_like_text(data: bytes) -> bool:
    sample = data[:8192]
    if sample.startswith((b"\xef\xbb\xbf", b"\xff\xfe", b"\xfe\xff")):
        return True
    if b"\x00" in sample:
        return False
    try:
        sample.decode("utf-8")
        return True
    except UnicodeDecodeError as exc:
        # a multi-byte character cut at the sample edge is still text
        if exc.start >= len(sample) - 4:
            return True
    try:
        decoded = sample.decode("cp1252")
    except UnicodeDecodeError:
        return False
    printable = sum(1 for ch in decoded if ch.isprintable() or ch in "\r\n\t")
    return printable / max(1, len(decoded)) > 0.95


_EXPECTED_SIGNATURE: Dict[str, Tuple[str, ...]] = {
    "pdf": ("pdf",),
    "docx": ("ooxml_docx",), "google_doc": ("ooxml_docx",),
    "xlsx": ("ooxml_xlsx",), "google_sheet": ("ooxml_xlsx",),
    "pptx": ("ooxml_pptx",), "google_slides": ("ooxml_pptx",),
    "xls": ("ole",),
    "rtf": ("rtf",),
    "zip": ("zip", "ooxml_docx", "ooxml_xlsx", "ooxml_pptx"),
    "image": ("png", "jpeg", "webp", "tiff", "gif", "bmp"),
    "txt": ("text",), "md": ("text",), "csv": ("text",), "json": ("text",), "xml": ("text",), "html": ("text",),
}

_SNIFF_TO_FORMAT = {
    "pdf": "pdf", "ooxml_docx": "docx", "ooxml_xlsx": "xlsx", "ooxml_pptx": "pptx", "zip": "zip", "rtf": "rtf",
    "png": "image", "jpeg": "image", "webp": "image", "tiff": "image", "gif": "image", "bmp": "image",
}


def verify_signature(fmt: str, data: bytes) -> Tuple[Optional[str], Optional[Tuple[str, str]]]:
    """
    Check the bytes against the declared format.
    Returns (format_to_use, None) when acceptable, or (None, (status, reason)) when the file must not be parsed.
    A file whose real content is a different *supported document* (e.g. a PDF named .docx) is parsed as what it is.
    """
    kind = sniff(data)
    if kind == "empty":
        return None, (ST_EMPTY, "The file is empty (0 bytes).")
    if kind == "executable":
        return None, (ST_UNSUPPORTED, "The file content is an executable or script; it was not opened.")
    if kind == "ole" and fmt in ("docx", "xlsx", "pptx", "google_doc", "google_sheet", "google_slides"):
        return None, (ST_ENCRYPTED, "The Office file is encrypted or password protected. Remove the password and re-run.")
    if kind == "ole" and fmt in ("doc", "ppt"):
        return None, (ST_UNSUPPORTED, UNSUPPORTED_REASONS[fmt])
    expected = _EXPECTED_SIGNATURE.get(fmt, ())
    if kind in expected:
        return fmt, None
    actual = _SNIFF_TO_FORMAT.get(kind)
    if actual and actual != fmt:
        return actual, None
    if kind == "text" and fmt in TEXT_FORMATS:
        return fmt, None
    return None, (ST_ERROR, f"The file content ({kind}) does not match its declared type ({fmt}); it was not parsed.")


# --------------------------------------------------------------------------------------------------------------------
# Result model
# --------------------------------------------------------------------------------------------------------------------
class ExtractionTimeout(Exception):
    pass


class Deadline:
    """Cooperative time limit checked inside parsing loops (pages, rows, slides, archive members)."""

    def __init__(self, seconds: Optional[float]):
        self.until = time.monotonic() + float(seconds) if seconds else None

    def check(self) -> None:
        if self.until is not None and time.monotonic() > self.until:
            raise ExtractionTimeout("processing time limit reached")


@dataclass
class Segment:
    text: str
    ref: Dict[str, Any]
    kind: str = "text"      # page, paragraph, heading, table_row, slide, notes, header, footer, ocr, cell


@dataclass
class ExtractionResult:
    format: str
    status: str = ST_OK
    segments: List[Segment] = field(default_factory=list)
    records: List[Dict[str, Any]] = field(default_factory=list)   # {"data": {...}, "ref": {...}}
    images: List[Dict[str, Any]] = field(default_factory=list)    # {"name", "ref", "mime"}
    warnings: List[str] = field(default_factory=list)
    error: Optional[str] = None
    meta: Dict[str, Any] = field(default_factory=dict)
    children: List[Tuple[str, "ExtractionResult"]] = field(default_factory=list)  # (archive member name, result)

    @property
    def usable(self) -> bool:
        return self.status in USABLE_STATUSES

    def text(self, limit: Optional[int] = None) -> str:
        out = "\n".join(s.text for s in self.segments if s.text)
        return out[:limit] if limit else out

    def summary(self) -> Dict[str, Any]:
        return {
            "format": self.format,
            "status": self.status,
            "segments": len(self.segments),
            "records": len(self.records),
            "images": len(self.images),
            "warnings": list(self.warnings)[:20],
            "error": self.error,
            "meta": dict(self.meta),
            "archive_members": [
                {"name": name, "format": child.format, "status": child.status, "error": child.error,
                 "records": len(child.records), "segments": len(child.segments)}
                for name, child in self.children
            ],
        }


def _fail(fmt: str, status: str, error: str) -> ExtractionResult:
    return ExtractionResult(format=fmt, status=status, error=error)


def _finish(res: ExtractionResult) -> ExtractionResult:
    if res.status == ST_OK and not res.segments and not res.records:
        res.status = ST_EMPTY
        res.error = res.error or "No readable text or rows were found in the file."
    return res


def _max_rows() -> int:
    return int(_cfg("DRIVE_INGEST_MAX_ROWS", 5000))


def _clean_cell(value: Any) -> Any:
    if value is None:
        return None
    if isinstance(value, float) and value != value:  # NaN
        return None
    if isinstance(value, str):
        value = value.strip()
        return value or None
    if hasattr(value, "isoformat"):
        return value.isoformat()
    return value


def _header_row(rows: List[List[Any]]) -> Tuple[Optional[int], List[str]]:
    """First row with at least two non-empty text cells, within the first 10 rows."""
    for idx, row in enumerate(rows[:10]):
        cells = [_clean_cell(c) for c in row]
        texts = [c for c in cells if isinstance(c, str)]
        if len(texts) >= 2 and len(texts) >= len([c for c in cells if c is not None]) * 0.6:
            headers = []
            for pos, c in enumerate(cells):
                name = str(c).strip() if c is not None else ""
                headers.append(name or f"column_{pos + 1}")
            return idx, headers
    return None, []


def _rows_to_records(rows: List[List[Any]], base_ref: Dict[str, Any], row_offset: int = 1,
                     deadline: Optional[Deadline] = None) -> Tuple[List[Dict[str, Any]], List[Segment], bool]:
    """Tabular rows -> records keyed by header, plus one text segment per row. Returns (records, segments, truncated)."""
    records: List[Dict[str, Any]] = []
    segments: List[Segment] = []
    header_idx, headers = _header_row(rows)
    limit = _max_rows()
    truncated = False
    start = (header_idx + 1) if header_idx is not None else 0
    if header_idx is not None:
        segments.append(Segment(" | ".join(headers), {**base_ref, "row": header_idx + row_offset}, "table_header"))
    for idx in range(start, len(rows)):
        if deadline and idx % 200 == 0:
            deadline.check()
        cells = [_clean_cell(c) for c in rows[idx]]
        if not any(c is not None for c in cells):
            continue
        if len(records) >= limit:
            truncated = True
            break
        ref = {**base_ref, "row": idx + row_offset}
        if headers:
            data = {}
            for pos, value in enumerate(cells):
                if value is None:
                    continue
                key = headers[pos] if pos < len(headers) else f"column_{pos + 1}"
                data[key] = value
            if data:
                records.append({"data": data, "ref": ref})
        segments.append(Segment(" | ".join("" if c is None else str(c) for c in cells), ref, "table_row"))
    return records, segments, truncated


# --------------------------------------------------------------------------------------------------------------------
# PDF
# --------------------------------------------------------------------------------------------------------------------
def extract_pdf(data: bytes, deadline: Deadline, ocr: Optional[Callable[[bytes, str], str]] = None) -> ExtractionResult:
    import pypdf

    res = ExtractionResult(format="pdf")
    try:
        reader = pypdf.PdfReader(io.BytesIO(data))
    except Exception as exc:
        return _fail("pdf", ST_ERROR, f"The PDF could not be opened ({type(exc).__name__}); it may be corrupted.")
    if reader.is_encrypted:
        try:
            # Owner-password-only PDFs (printing restrictions) open with an empty user password.
            if not reader.decrypt(""):
                return _fail("pdf", ST_ENCRYPTED, "The PDF is password protected. Remove the password and re-run.")
        except Exception:
            return _fail("pdf", ST_ENCRYPTED, "The PDF is password protected. Remove the password and re-run.")
    try:
        total = len(reader.pages)
    except Exception as exc:
        return _fail("pdf", ST_ERROR, f"The PDF page tree is damaged ({type(exc).__name__}).")
    max_pages = int(_cfg("DRIVE_INGEST_MAX_PDF_PAGES", 300))
    ocr_budget = int(_cfg("DRIVE_INGEST_OCR_MAX_PAGES", 10))
    text_pages = 0
    ocr_pages = 0
    image_only_pages: List[int] = []
    ocr_filled: List[int] = []
    for i in range(min(total, max_pages)):
        deadline.check()
        page_no = i + 1
        try:
            page = reader.pages[i]
            text = (page.extract_text() or "").strip()
        except Exception as exc:
            res.warnings.append(f"page {page_no}: text could not be read ({type(exc).__name__})")
            continue
        page_images = []
        try:
            page_images = list(page.images)
        except Exception:
            page_images = []
        for img in page_images[:20]:
            res.images.append({"name": safe_filename(getattr(img, "name", "") or f"image_p{page_no}"),
                               "ref": {"page": page_no}, "mime": None})
        if text:
            text_pages += 1
            res.segments.append(Segment(text, {"page": page_no}, "page"))
            continue
        image_only_pages.append(page_no)
        if ocr and page_images and ocr_pages < ocr_budget:
            biggest = max(page_images, key=lambda im: len(getattr(im, "data", b"") or b""))
            try:
                img_bytes = biggest.data
                ocr_text = (ocr(img_bytes, "page") or "").strip()
            except Exception as exc:
                res.warnings.append(f"page {page_no}: OCR failed ({type(exc).__name__})")
                ocr_text = ""
            ocr_pages += 1
            if ocr_text:
                ocr_filled.append(page_no)
                res.segments.append(Segment(ocr_text, {"page": page_no, "ocr": True}, "ocr"))
    if total > max_pages:
        res.warnings.append(f"only the first {max_pages} of {total} pages were read")
    res.meta.update({"pages": total, "text_pages": text_pages, "ocr_pages": ocr_pages,
                     "image_only_pages": image_only_pages[:50]})
    if total == 0:
        res.status, res.error = ST_EMPTY, "The PDF has no pages."
    elif not res.segments:
        res.status = ST_NEEDS_OCR
        res.error = ("The PDF is scanned or image-only and has no text layer. "
                     + ("OCR returned no text." if ocr else "OCR is not enabled (DRIVE_INGEST_OCR_ENABLED), so nothing was read."))
    elif len(image_only_pages) > len(ocr_filled):
        res.status = ST_PARTIAL
        res.warnings.append(f"{len(image_only_pages) - len(ocr_filled)} page(s) have no text layer"
                            + (" and OCR did not read them" if ocr else " and OCR is off"))
    if total > max_pages and res.status == ST_OK:
        res.status = ST_PARTIAL
    return res


# --------------------------------------------------------------------------------------------------------------------
# Word (.docx / Google Docs export)
# --------------------------------------------------------------------------------------------------------------------
def extract_docx(data: bytes, deadline: Deadline, fmt: str = "docx") -> ExtractionResult:
    try:
        import docx
        from docx.table import Table
        from docx.text.paragraph import Paragraph
    except Exception as exc:  # pragma: no cover - dependency is in requirements
        return _fail(fmt, ST_ERROR, f"Word parser unavailable ({exc})")
    res = ExtractionResult(format=fmt)
    try:
        document = docx.Document(io.BytesIO(data))
    except Exception as exc:
        return _fail(fmt, ST_ERROR, f"The Word document could not be opened ({type(exc).__name__}); it may be corrupted.")

    heading = None
    para_no = 0
    table_no = 0
    body = document.element.body
    for child in body.iterchildren():
        deadline.check()
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            para = Paragraph(child, document)
            text = (para.text or "").strip()
            if not text:
                continue
            para_no += 1
            style = (getattr(para.style, "name", "") or "").lower() if para.style is not None else ""
            is_heading = style.startswith("heading") or style == "title"
            if is_heading:
                heading = text[:200]
            ref = {"paragraph": para_no}
            if heading and not is_heading:
                ref["section"] = heading
            res.segments.append(Segment(text, ref, "heading" if is_heading else "paragraph"))
        elif tag == "tbl":
            table_no += 1
            table = Table(child, document)
            rows: List[List[Any]] = []
            for row in table.rows:
                cells = []
                seen = set()
                for cell in row.cells:
                    # merged cells repeat the same object; keep one copy
                    if id(cell._tc) in seen:
                        continue
                    seen.add(id(cell._tc))
                    cells.append((cell.text or "").strip())
                rows.append(cells)
            base = {"table": table_no}
            if heading:
                base["section"] = heading
            records, segments, truncated = _rows_to_records(rows, base, deadline=deadline)
            res.records.extend(records)
            res.segments.extend(segments)
            if truncated:
                res.warnings.append(f"table {table_no}: row limit reached")
    for s_idx, section in enumerate(document.sections, start=1):
        for part_name, part in (("header", section.header), ("footer", section.footer)):
            try:
                if part is None or part.is_linked_to_previous and s_idx > 1:
                    continue
                text = "\n".join(p.text.strip() for p in part.paragraphs if p.text and p.text.strip())
            except Exception:
                text = ""
            if text:
                res.segments.append(Segment(text, {part_name: s_idx}, part_name))
    try:
        for rel in document.part.rels.values():
            if "image" in rel.reltype:
                res.images.append({"name": safe_filename(os.path.basename(getattr(rel, "target_ref", "") or "image")),
                                   "ref": {"document": True}, "mime": getattr(getattr(rel, "target_part", None), "content_type", None)})
    except Exception:
        pass
    res.meta.update({"paragraphs": para_no, "tables": table_no})
    return _finish(res)


# --------------------------------------------------------------------------------------------------------------------
# Spreadsheets (.xlsx / Google Sheets export / .xls / .csv)
# --------------------------------------------------------------------------------------------------------------------
def extract_xlsx(data: bytes, deadline: Deadline, fmt: str = "xlsx") -> ExtractionResult:
    import openpyxl

    res = ExtractionResult(format=fmt)
    try:
        # read_only streams rows; data_only returns cached values — formulas are never evaluated.
        wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:
        return _fail(fmt, ST_ERROR, f"The workbook could not be opened ({type(exc).__name__}); it may be corrupted.")
    sheets = []
    limit = _max_rows()
    try:
        for ws in wb.worksheets:
            deadline.check()
            rows: List[List[Any]] = []
            for r_idx, row in enumerate(ws.iter_rows(values_only=True)):
                if r_idx % 500 == 0:
                    deadline.check()
                rows.append(list(row))
                if len(rows) > limit + 10:
                    res.warnings.append(f"sheet '{ws.title}': row limit {limit} reached")
                    break
            records, segments, truncated = _rows_to_records(rows, {"sheet": ws.title}, deadline=deadline)
            res.records.extend(records)
            res.segments.extend(segments)
            if truncated:
                res.warnings.append(f"sheet '{ws.title}': row limit {limit} reached")
            sheets.append({"name": ws.title, "rows": len(rows), "records": len(records)})
    finally:
        try:
            wb.close()
        except Exception:
            pass
    res.meta["sheets"] = sheets
    if any("row limit" in w for w in res.warnings):
        res.status = ST_PARTIAL
    return _finish(res)


def extract_xls(data: bytes, deadline: Deadline) -> ExtractionResult:
    try:
        import xlrd
    except Exception:
        return _fail("xls", ST_UNSUPPORTED, "Legacy .xls support needs the 'xlrd' package, which is not installed.")
    res = ExtractionResult(format="xls")
    try:
        book = xlrd.open_workbook(file_contents=data, on_demand=True)
    except Exception as exc:
        msg = str(exc).lower()
        if "encrypt" in msg or "password" in msg:
            return _fail("xls", ST_ENCRYPTED, "The workbook is password protected. Remove the password and re-run.")
        return _fail("xls", ST_ERROR, f"The .xls workbook could not be opened ({type(exc).__name__}).")
    sheets = []
    try:
        for sheet in book.sheets():
            deadline.check()
            rows = [sheet.row_values(i) for i in range(min(sheet.nrows, _max_rows() + 10))]
            records, segments, truncated = _rows_to_records(rows, {"sheet": sheet.name}, deadline=deadline)
            res.records.extend(records)
            res.segments.extend(segments)
            if truncated or sheet.nrows > _max_rows() + 10:
                res.warnings.append(f"sheet '{sheet.name}': row limit reached")
                res.status = ST_PARTIAL
            sheets.append({"name": sheet.name, "rows": sheet.nrows, "records": len(records)})
    finally:
        try:
            book.release_resources()
        except Exception:
            pass
    res.meta["sheets"] = sheets
    return _finish(res)


def decode_text(data: bytes) -> str:
    if data[:2] in (b"\xff\xfe", b"\xfe\xff"):
        try:
            return data.decode("utf-16")
        except UnicodeDecodeError:
            pass
    try:
        return data.decode("utf-8-sig")
    except UnicodeDecodeError:
        return data.decode("cp1252", errors="replace")


def extract_csv(data: bytes, deadline: Deadline) -> ExtractionResult:
    res = ExtractionResult(format="csv")
    text = decode_text(data)
    sample = text[:8192]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    csv.field_size_limit(1_000_000)
    try:
        rows = []
        for idx, row in enumerate(csv.reader(io.StringIO(text), dialect)):
            if idx % 500 == 0:
                deadline.check()
            rows.append(row)
            if len(rows) > _max_rows() + 10:
                res.warnings.append(f"row limit {_max_rows()} reached")
                res.status = ST_PARTIAL
                break
    except csv.Error as exc:
        return _fail("csv", ST_ERROR, f"The CSV is malformed ({exc}).")
    records, segments, truncated = _rows_to_records(rows, {"sheet": "csv"}, deadline=deadline)
    res.records, res.segments = records, segments
    if truncated:
        res.status = ST_PARTIAL
    res.meta["rows"] = len(rows)
    return _finish(res)


# --------------------------------------------------------------------------------------------------------------------
# PowerPoint (.pptx / Google Slides export)
# --------------------------------------------------------------------------------------------------------------------
def extract_pptx(data: bytes, deadline: Deadline, fmt: str = "pptx") -> ExtractionResult:
    try:
        from pptx import Presentation
    except Exception:
        return _fail(fmt, ST_UNSUPPORTED, "PowerPoint support needs the 'python-pptx' package, which is not installed.")
    res = ExtractionResult(format=fmt)
    try:
        prs = Presentation(io.BytesIO(data))
    except Exception as exc:
        return _fail(fmt, ST_ERROR, f"The presentation could not be opened ({type(exc).__name__}); it may be corrupted.")
    slides = 0
    for s_idx, slide in enumerate(prs.slides, start=1):
        deadline.check()
        slides += 1
        title = None
        try:
            if slide.shapes.title is not None and slide.shapes.title.has_text_frame:
                title = (slide.shapes.title.text_frame.text or "").strip() or None
        except Exception:
            title = None
        base = {"slide": s_idx}
        if title:
            base["slide_title"] = title[:200]
            res.segments.append(Segment(title, dict(base), "heading"))
        table_no = 0
        for shape in _iter_shapes(slide.shapes):
            if shape is getattr(slide.shapes, "title", None):
                continue
            if getattr(shape, "has_table", False) and shape.has_table:
                table_no += 1
                rows = [[(cell.text or "").strip() for cell in row.cells] for row in shape.table.rows]
                records, segments, _ = _rows_to_records(rows, {**base, "table": table_no}, deadline=deadline)
                res.records.extend(records)
                res.segments.extend(segments)
            elif getattr(shape, "has_text_frame", False) and shape.has_text_frame:
                text = (shape.text_frame.text or "").strip()
                if text and text != title:
                    res.segments.append(Segment(text, dict(base), "slide"))
            if shape.__class__.__name__ == "Picture":
                res.images.append({"name": safe_filename(getattr(shape, "name", "") or f"slide{s_idx}_image"),
                                   "ref": dict(base), "mime": getattr(getattr(shape, "image", None), "content_type", None)})
        try:
            if slide.has_notes_slide:
                notes = (slide.notes_slide.notes_text_frame.text or "").strip()
                if notes:
                    res.segments.append(Segment(notes, {**base, "notes": True}, "notes"))
        except Exception:
            pass
    res.meta["slides"] = slides
    return _finish(res)


def _iter_shapes(shapes):
    for shape in shapes:
        if shape.__class__.__name__ == "GroupShape":
            yield from _iter_shapes(shape.shapes)
        else:
            yield shape


# --------------------------------------------------------------------------------------------------------------------
# Text, Markdown, RTF
# --------------------------------------------------------------------------------------------------------------------
def _blocks(text: str) -> List[str]:
    return [b.strip() for b in re.split(r"\n\s*\n", text) if b.strip()]


def extract_text(data: bytes, deadline: Deadline, fmt: str = "txt") -> ExtractionResult:
    res = ExtractionResult(format=fmt)
    text = decode_text(data)
    heading = None
    for idx, block in enumerate(_blocks(text), start=1):
        if idx % 200 == 0:
            deadline.check()
        ref: Dict[str, Any] = {"block": idx}
        kind = "paragraph"
        if fmt == "md":
            first = block.splitlines()[0]
            m = re.match(r"^#{1,6}\s+(.*)", first)
            if m:
                heading = m.group(1).strip()[:200]
                kind = "heading"
            # markdown pipe tables become records
            lines = [ln for ln in block.splitlines() if ln.strip().startswith("|")]
            if len(lines) >= 2 and re.match(r"^\|?\s*:?-{3,}", lines[1].strip().strip("|").strip() or "---"):
                rows = [[c.strip() for c in ln.strip().strip("|").split("|")] for ln in lines if not re.match(r"^\|?[\s:|-]+$", ln.strip())]
                records, segments, _ = _rows_to_records(rows, {"block": idx, **({"section": heading} if heading else {})})
                res.records.extend(records)
        if heading and kind != "heading":
            ref["section"] = heading
        res.segments.append(Segment(block, ref, kind))
    return _finish(res)


_RTF_DESTINATIONS = {
    "aftncn", "aftnsep", "aftnsepc", "annotation", "atnauthor", "atndate", "atnicn", "atnid", "atnparent", "atnref",
    "atntime", "atrfend", "atrfstart", "author", "background", "bkmkend", "bkmkstart", "blipuid", "buptim", "category",
    "colorschememapping", "colortbl", "comment", "company", "creatim", "datafield", "datastore", "defchp", "defpap",
    "do", "doccomm", "docvar", "dptxbxtext", "ebcend", "ebcstart", "factoidname", "falt", "fchars", "ffdeftext",
    "ffentrymcr", "ffexitmcr", "ffformat", "ffhelptext", "ffl", "ffname", "ffstattext", "field", "file", "filetbl",
    "fldinst", "fldtype", "fname", "fontemb", "fontfile", "fonttbl", "footer", "footerf", "footerl", "footerr",
    "footnote", "formfield", "ftncn", "ftnsep", "ftnsepc", "g", "generator", "gridtbl", "header", "headerf", "headerl",
    "headerr", "hl", "hlfr", "hlinkbase", "hlloc", "hlsrc", "hsv", "htmltag", "info", "keycode", "keywords",
    "latentstyles", "lchars", "levelnumbers", "leveltext", "lfolevel", "linkval", "list", "listlevel", "listname",
    "listoverride", "listoverridetable", "listpicture", "liststylename", "listtable", "listtext", "lsdlockedexcept",
    "macc", "maccPr", "mailmerge", "maln", "malnScr", "manager", "margPr", "mbar", "mbarPr", "mbaseJc", "mbegChr",
    "mborderBox", "mborderBoxPr", "mbox", "mboxPr", "mchr", "mcount", "mctrlPr", "md", "mdeg", "mdegHide", "mden",
    "mdiff", "mdPr", "me", "mendChr", "meqArr", "meqArrPr", "mf", "mfName", "mfPr", "mfunc", "mfuncPr", "mgroupChr",
    "mgroupChrPr", "mgrow", "mhideBot", "mhideLeft", "mhideRight", "mhideTop", "mhtmltag", "mlim", "mlimloc",
    "mlimlow", "mlimlowPr", "mlimupp", "mlimuppPr", "mm", "mmaddfieldname", "mmath", "mmathPict", "mmathPr",
    "mmaxdist", "mmc", "mmcJc", "mmconnectstr", "mmconnectstrdata", "mmcPr", "mmcs", "mmdatasource", "mmheadersource",
    "mmmailsubject", "mmodso", "mmodsofilter", "mmodsofldmpdata", "mmodsomappedname", "mmodsoname", "mmodsorecipdata",
    "mmodsosort", "mmodsosrc", "mmodsotable", "mmodsoudl", "mmodsoudldata", "mmodsouniquetag", "mmPr", "mmquery",
    "mmr", "mnary", "mnaryPr", "mnoBreak", "mnum", "mobjDist", "moMath", "moMathPara", "moMathParaPr", "mopEmu",
    "mphant", "mphantPr", "mplcHide", "mpos", "mr", "mrad", "mradPr", "mrPr", "msepChr", "mshow", "mshp", "msPre",
    "msPrePr", "msSub", "msSubPr", "msSubSup", "msSubSupPr", "msSup", "msSupPr", "mstrikeBLTR", "mstrikeH",
    "mstrikeTLBR", "mstrikeV", "msub", "msubHide", "msup", "msupHide", "mtransp", "mtype", "mvertJc", "mvfmf", "mvfml",
    "mvtof", "mvtol", "mzeroAsc", "mzeroDesc", "mzeroWid", "nesttableprops", "nextfile", "nonesttables", "objalias",
    "objclass", "objdata", "object", "objname", "objsect", "objtime", "oldcprops", "oldpprops", "oldsprops",
    "oldtprops", "oleclsid", "operator", "panose", "password", "passwordhash", "pgp", "pgptbl", "picprop", "pict",
    "pn", "pnseclvl", "pntext", "pntxta", "pntxtb", "printim", "private", "propname", "protend", "protstart",
    "protusertbl", "pxe", "result", "revtbl", "revtim", "rsidtbl", "rxe", "shp", "shpgrp", "shpinst", "shppict",
    "shprslt", "shptxt", "sn", "sp", "staticval", "stylesheet", "subject", "sv", "svb", "tc", "template", "themedata",
    "title", "txe", "ud", "upr", "userprops", "wgrffmtfilter", "windowcaption", "writereservation",
    "writereservhash", "xe", "xform", "xmlattrname", "xmlattrvalue", "xmlclose", "xmlname", "xmlnstbl", "xmlopen",
}
_RTF_SPECIAL = {"par": "\n", "sect": "\n\n", "page": "\n\n", "line": "\n", "tab": "\t", "emdash": "\u2014",
                "endash": "\u2013", "emspace": " ", "enspace": " ", "qmspace": " ", "bullet": "\u2022",
                "lquote": "\u2018", "rquote": "\u2019", "ldblquote": "\u201c", "rdblquote": "\u201d", "row": "\n",
                "cell": " | "}
_RTF_PATTERN = re.compile(r"\\([a-z]{1,32})(-?\d{1,10})?[ ]?|\\'([0-9a-f]{2})|\\([^a-z])|([{}])|[\r\n]+|(.)", re.I)


def rtf_to_text(rtf: str) -> str:
    """Plain text from RTF. Embedded objects, pictures, fields' instructions and metadata groups are dropped."""
    stack: List[Tuple[int, bool]] = []
    ignorable = False
    ucskip = 1
    curskip = 0
    out: List[str] = []
    for match in _RTF_PATTERN.finditer(rtf):
        word, arg, hexcode, char, brace, tchar = match.groups()
        if brace:
            curskip = 0
            if brace == "{":
                stack.append((ucskip, ignorable))
            elif stack:
                ucskip, ignorable = stack.pop()
        elif char:
            curskip = 0
            if char == "~":
                if not ignorable:
                    out.append("\xa0")
            elif char in "{}\\":
                if not ignorable:
                    out.append(char)
            elif char == "*":
                ignorable = True
        elif word:
            curskip = 0
            if word in _RTF_DESTINATIONS:
                ignorable = True
            elif ignorable:
                pass
            elif word in _RTF_SPECIAL:
                out.append(_RTF_SPECIAL[word])
            elif word == "uc":
                ucskip = int(arg or 1)
            elif word == "u":
                c = int(arg or 0)
                if c < 0:
                    c += 0x10000
                try:
                    out.append(chr(c))
                except ValueError:
                    pass
                curskip = ucskip
        elif hexcode:
            if curskip > 0:
                curskip -= 1
            elif not ignorable:
                out.append(bytes([int(hexcode, 16)]).decode("cp1252", errors="replace"))
        elif tchar:
            if curskip > 0:
                curskip -= 1
            elif not ignorable:
                out.append(tchar)
    return "".join(out)


def extract_rtf(data: bytes, deadline: Deadline) -> ExtractionResult:
    try:
        text = rtf_to_text(data.decode("latin-1"))
    except Exception as exc:
        return _fail("rtf", ST_ERROR, f"The RTF file could not be read ({type(exc).__name__}).")
    deadline.check()
    res = extract_text(text.encode("utf-8"), deadline, fmt="rtf")
    res.format = "rtf"
    return res


# --------------------------------------------------------------------------------------------------------------------
# HTML
# --------------------------------------------------------------------------------------------------------------------
class _HtmlExtractor(HTMLParser):
    _SKIP = {"script", "style", "noscript", "template", "iframe", "object", "embed", "svg", "canvas", "head"}
    _BLOCK = {"p", "div", "li", "h1", "h2", "h3", "h4", "h5", "h6", "section", "article", "br", "tr", "td", "th",
              "dd", "dt", "blockquote", "pre", "header", "footer", "main", "aside", "table", "caption"}

    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.skip_depth = 0
        self.title = ""
        self.in_title = False
        self.meta: Dict[str, str] = {}
        self.ldjson: List[str] = []
        self.in_ldjson = False
        self._ld_buf: List[str] = []
        self.blocks: List[Tuple[str, str, Optional[str]]] = []   # (kind, text, heading)
        self._buf: List[str] = []
        self._kind = "paragraph"
        self.heading: Optional[str] = None
        self.tables: List[List[List[str]]] = []
        self._table: Optional[List[List[str]]] = None
        self._row: Optional[List[str]] = None
        self._cell: Optional[List[str]] = None
        self.images: List[str] = []

    def _flush(self):
        text = re.sub(r"\s+", " ", "".join(self._buf)).strip()
        self._buf = []
        if text:
            if self._kind == "heading":
                self.heading = text[:200]
            self.blocks.append((self._kind, text, self.heading if self._kind != "heading" else None))
        self._kind = "paragraph"

    def handle_starttag(self, tag, attrs):
        a = dict(attrs)
        if tag == "script" and (a.get("type") or "").lower() == "application/ld+json":
            self.in_ldjson = True
            self._ld_buf = []
            return
        if tag in self._SKIP and tag != "head":
            self.skip_depth += 1
            return
        if tag == "title":
            self.in_title = True
        if tag == "meta":
            key = (a.get("property") or a.get("name") or "").lower()
            if key and a.get("content"):
                self.meta[key] = a["content"][:500]
        if tag == "img" and a.get("src") and len(self.images) < 50:
            self.images.append(a.get("alt") or a.get("src"))
        if self.skip_depth:
            return
        if tag == "table":
            self._flush()
            self._table = []
        elif tag == "tr" and self._table is not None:
            self._row = []
        elif tag in ("td", "th") and self._row is not None:
            self._cell = []
        elif tag in self._BLOCK:
            self._flush()
            if re.match(r"h[1-6]$", tag):
                self._kind = "heading"

    def handle_endtag(self, tag):
        if tag == "script" and self.in_ldjson:
            self.in_ldjson = False
            self.ldjson.append("".join(self._ld_buf))
            return
        if tag in self._SKIP and tag != "head":
            self.skip_depth = max(0, self.skip_depth - 1)
            return
        if tag == "title":
            self.in_title = False
        if self.skip_depth:
            return
        if tag in ("td", "th") and self._cell is not None and self._row is not None:
            self._row.append(re.sub(r"\s+", " ", "".join(self._cell)).strip())
            self._cell = None
        elif tag == "tr" and self._row is not None and self._table is not None:
            if any(self._row):
                self._table.append(self._row)
            self._row = None
        elif tag == "table" and self._table is not None:
            if self._table:
                self.tables.append(self._table)
            self._table = None
        elif tag in self._BLOCK:
            self._flush()

    def handle_data(self, data):
        if self.in_ldjson:
            self._ld_buf.append(data)
            return
        if self.in_title:
            self.title += data
            return
        if self.skip_depth:
            return
        if self._cell is not None:
            self._cell.append(data)
        elif self._table is None:
            self._buf.append(data)


def extract_html(data: bytes, deadline: Deadline) -> ExtractionResult:
    res = ExtractionResult(format="html")
    parser = _HtmlExtractor()
    try:
        parser.feed(decode_text(data))
        parser.close()
    except Exception as exc:
        return _fail("html", ST_ERROR, f"The HTML could not be parsed ({type(exc).__name__}).")
    deadline.check()
    if parser.title.strip():
        res.segments.append(Segment(parser.title.strip(), {"element": "title"}, "heading"))
    for key in ("description", "og:description", "og:title"):
        if parser.meta.get(key):
            res.segments.append(Segment(parser.meta[key], {"element": f"meta:{key}"}, "meta"))
    for idx, (kind, text, heading) in enumerate(parser.blocks, start=1):
        ref: Dict[str, Any] = {"block": idx}
        if heading:
            ref["section"] = heading
        res.segments.append(Segment(text, ref, kind))
    for t_idx, rows in enumerate(parser.tables, start=1):
        records, segments, _ = _rows_to_records(rows, {"table": t_idx}, deadline=deadline)
        res.records.extend(records)
        res.segments.extend(segments)
    for b_idx, raw in enumerate(parser.ldjson, start=1):
        try:
            obj = json.loads(raw)
        except Exception:
            continue
        for item in (obj if isinstance(obj, list) else [obj]):
            if isinstance(item, dict):
                res.records.append({"data": _flatten(item), "ref": {"json_ld": b_idx}})
    res.images = [{"name": safe_filename(n), "ref": {"element": "img"}, "mime": None} for n in parser.images]
    res.meta.update({"title": parser.title.strip()[:200], "tables": len(parser.tables)})
    return _finish(res)


# --------------------------------------------------------------------------------------------------------------------
# JSON
# --------------------------------------------------------------------------------------------------------------------
_RECORD_KEYS = ("properties", "listings", "items", "data", "results", "records", "projects", "units", "rows")


def _flatten(obj: Dict[str, Any], prefix: str = "", depth: int = 0) -> Dict[str, Any]:
    """Nested dicts become dotted keys (address.city). Lists of scalars are kept; lists of objects are summarised."""
    out: Dict[str, Any] = {}
    for key, value in obj.items():
        name = f"{prefix}{key}"
        if isinstance(value, dict) and depth < 4:
            out.update(_flatten(value, f"{name}.", depth + 1))
        elif isinstance(value, list):
            if all(not isinstance(v, (dict, list)) for v in value):
                out[name] = [v for v in value if v is not None][:100]
            else:
                out[name] = [_flatten(v, "", depth + 1) if isinstance(v, dict) else v for v in value[:20]]
        else:
            out[name] = value
    return out


def _json_depth_ok(text: str, limit: int = 64) -> bool:
    depth = 0
    in_str = False
    esc = False
    for ch in text:
        if in_str:
            if esc:
                esc = False
            elif ch == "\\":
                esc = True
            elif ch == '"':
                in_str = False
        elif ch == '"':
            in_str = True
        elif ch in "[{":
            depth += 1
            if depth > limit:
                return False
        elif ch in "]}":
            depth -= 1
    return True


def extract_json(data: bytes, deadline: Deadline) -> ExtractionResult:
    res = ExtractionResult(format="json")
    text = decode_text(data)
    if not _json_depth_ok(text):
        return _fail("json", ST_ERROR, "The JSON nests deeper than 64 levels and was not parsed.")
    try:
        obj = json.loads(text)
    except Exception as exc:
        return _fail("json", ST_ERROR, f"The JSON is malformed ({type(exc).__name__}: {str(exc)[:80]}).")
    deadline.check()
    path = "$"
    items: Optional[List[Any]] = None
    if isinstance(obj, list):
        items = obj
    elif isinstance(obj, dict):
        for key in _RECORD_KEYS:
            if isinstance(obj.get(key), list) and any(isinstance(v, dict) for v in obj[key]):
                items, path = obj[key], f"$.{key}"
                break
        if items is None:
            # a single dict may still wrap a list under an unknown key
            lists = [(k, v) for k, v in obj.items() if isinstance(v, list) and v and all(isinstance(i, dict) for i in v)]
            if len(lists) == 1:
                items, path = lists[0][1], f"$.{lists[0][0]}"
    if items is None:
        if isinstance(obj, dict):
            res.records.append({"data": _flatten(obj), "ref": {"json_path": "$"}})
    else:
        for idx, item in enumerate(items):
            if idx % 200 == 0:
                deadline.check()
            if len(res.records) >= _max_rows():
                res.warnings.append(f"record limit {_max_rows()} reached")
                res.status = ST_PARTIAL
                break
            if isinstance(item, dict):
                res.records.append({"data": _flatten(item), "ref": {"json_path": f"{path}[{idx}]"}})
    for rec in res.records:
        flat = "; ".join(f"{k}: {v}" for k, v in rec["data"].items() if not isinstance(v, list))
        res.segments.append(Segment(flat[:4000], rec["ref"], "record"))
    return _finish(res)


# --------------------------------------------------------------------------------------------------------------------
# XML (defusedxml: no external entities, no entity expansion bombs, no DTD retrieval)
# --------------------------------------------------------------------------------------------------------------------
def _xml_local(tag: str) -> str:
    return tag.rsplit("}", 1)[-1] if isinstance(tag, str) else str(tag)


def _xml_to_dict(elem, depth: int = 0) -> Dict[str, Any]:
    out: Dict[str, Any] = {}
    for key, value in elem.attrib.items():
        out[f"@{_xml_local(key)}"] = value
    children = list(elem)
    if not children:
        text = (elem.text or "").strip()
        if text:
            out["#text"] = text
        return out
    for child in children:
        name = _xml_local(child.tag)
        if depth >= 6:
            continue
        grand = list(child)
        if not grand and not child.attrib:
            value: Any = (child.text or "").strip() or None
        elif not grand:
            # <price currency="THB">18500000</price> -> price: 18500000, price_currency: THB
            value = (child.text or "").strip() or None
            for attr, attr_val in child.attrib.items():
                out[f"{name}_{_xml_local(attr)}"] = attr_val
        else:
            value = _xml_to_dict(child, depth + 1)
            if set(value.keys()) == {"#text"}:
                value = value["#text"]
        if name in out:
            if not isinstance(out[name], list):
                out[name] = [out[name]]
            out[name].append(value)
        else:
            out[name] = value
    return out


def extract_xml(data: bytes, deadline: Deadline) -> ExtractionResult:
    try:
        from defusedxml import ElementTree as DET
        from defusedxml.common import DefusedXmlException
    except Exception:
        return _fail("xml", ST_UNSUPPORTED, "XML support needs the 'defusedxml' package, which is not installed.")
    res = ExtractionResult(format="xml")
    try:
        root = DET.fromstring(data, forbid_dtd=True)
    except DefusedXmlException as exc:
        return _fail("xml", ST_ERROR, f"The XML uses DTDs or entities, which are refused for safety ({type(exc).__name__}).")
    except Exception as exc:
        return _fail("xml", ST_ERROR, f"The XML is malformed ({type(exc).__name__}).")
    deadline.check()
    # The record element is the most repeated child tag under the root (or under the root's only child).
    container = root
    for _ in range(3):
        kids = list(container)
        if len(kids) == 1 and len(list(kids[0])) > 1:
            container = kids[0]
        else:
            break
    counts: Dict[str, int] = {}
    for child in container:
        counts[_xml_local(child.tag)] = counts.get(_xml_local(child.tag), 0) + 1
    record_tag = max(counts, key=counts.get) if counts else None
    base_path = "/" + _xml_local(root.tag) + ("" if container is root else "/" + _xml_local(container.tag))
    if record_tag and counts[record_tag] >= 1 and list(container):
        idx = 0
        for child in container:
            if _xml_local(child.tag) != record_tag:
                continue
            idx += 1
            if idx % 200 == 0:
                deadline.check()
            if len(res.records) >= _max_rows():
                res.warnings.append(f"record limit {_max_rows()} reached")
                res.status = ST_PARTIAL
                break
            data_dict = _flatten(_xml_to_dict(child))
            if data_dict:
                res.records.append({"data": data_dict, "ref": {"xml_path": f"{base_path}/{record_tag}[{idx}]"}})
    else:
        data_dict = _flatten(_xml_to_dict(root))
        if data_dict:
            res.records.append({"data": data_dict, "ref": {"xml_path": base_path}})
    for rec in res.records:
        flat = "; ".join(f"{k}: {v}" for k, v in rec["data"].items() if not isinstance(v, list))
        res.segments.append(Segment(flat[:4000], rec["ref"], "record"))
    return _finish(res)


# --------------------------------------------------------------------------------------------------------------------
# Images
# --------------------------------------------------------------------------------------------------------------------
def extract_image(data: bytes, deadline: Deadline, ocr: Optional[Callable[[bytes, str], str]] = None) -> ExtractionResult:
    from PIL import Image

    res = ExtractionResult(format="image")
    max_pixels = int(_cfg("DRIVE_INGEST_MAX_IMAGE_PIXELS", 60_000_000))
    try:
        with Image.open(io.BytesIO(data)) as img:
            width, height = img.size
            fmt = (img.format or "").lower()
            if width * height > max_pixels:
                return _fail("image", ST_TOO_LARGE,
                             f"The image is {width}x{height} pixels, above the {max_pixels:,}-pixel safety limit.")
            img.verify()
    except Image.DecompressionBombError:
        return _fail("image", ST_TOO_LARGE, "The image exceeds the decompression safety limit.")
    except Exception as exc:
        return _fail("image", ST_ERROR, f"The image could not be decoded ({type(exc).__name__}); it may be corrupted.")
    res.meta.update({"width": width, "height": height, "image_format": fmt})
    res.images.append({"name": "file", "ref": {"image": 1}, "mime": f"image/{fmt}" if fmt else None})
    if not ocr:
        res.status = ST_NEEDS_OCR
        res.error = "Images need OCR, which is not enabled (DRIVE_INGEST_OCR_ENABLED), so no text was read."
        return res
    deadline.check()
    try:
        text = (ocr(data, "image") or "").strip()
    except Exception as exc:
        return _fail("image", ST_ERROR, f"OCR failed ({type(exc).__name__}).")
    if not text:
        res.status = ST_EMPTY
        res.error = "OCR found no readable text in the image."
        return res
    res.segments.append(Segment(text, {"image": 1, "ocr": True}, "ocr"))
    res.meta["ocr"] = True
    return res


# --------------------------------------------------------------------------------------------------------------------
# ZIP archives
# --------------------------------------------------------------------------------------------------------------------
def extract_zip(data: bytes, deadline: Deadline, ocr: Optional[Callable[[bytes, str], str]] = None,
                depth: int = 0) -> ExtractionResult:
    res = ExtractionResult(format="zip")
    max_members = int(_cfg("DRIVE_INGEST_ZIP_MAX_MEMBERS", 200))
    max_total = int(_cfg("DRIVE_INGEST_ZIP_MAX_TOTAL_BYTES", 100 * 1024 * 1024))
    max_ratio = float(_cfg("DRIVE_INGEST_ZIP_MAX_RATIO", 100))
    per_member = max_file_bytes()
    try:
        zf = zipfile.ZipFile(io.BytesIO(data))
    except Exception as exc:
        return _fail("zip", ST_ERROR, f"The archive could not be opened ({type(exc).__name__}); it may be corrupted.")
    with zf:
        infos = [i for i in zf.infolist() if not i.is_dir()]
        if len(infos) > max_members:
            return _fail("zip", ST_TOO_LARGE, f"The archive holds {len(infos)} files; the limit is {max_members}.")
        declared_total = sum(i.file_size for i in infos)
        if declared_total > max_total:
            return _fail("zip", ST_TOO_LARGE,
                         f"The archive expands to {declared_total:,} bytes; the limit is {max_total:,} (possible ZIP bomb).")
        read_total = 0
        for info in infos:
            deadline.check()
            name = safe_filename(info.filename)
            raw_name = info.filename.replace("\\", "/")
            if raw_name.startswith("/") or ".." in raw_name.split("/"):
                res.warnings.append(f"{name}: unsafe path inside the archive; skipped")
                res.children.append((name, _fail("unsupported", ST_UNSUPPORTED, "Unsafe path (absolute or '..') inside the archive.")))
                continue
            if raw_name.startswith("__MACOSX/") or name.startswith("._") or name in (".DS_Store", "Thumbs.db"):
                continue
            if info.flag_bits & 0x1:
                res.children.append((name, _fail("unsupported", ST_ENCRYPTED, "The archive member is password protected.")))
                continue
            if info.compress_size and info.file_size / max(1, info.compress_size) > max_ratio:
                res.children.append((name, _fail("unsupported", ST_TOO_LARGE,
                                                 f"Compression ratio above {max_ratio:g}:1 (possible ZIP bomb); not extracted.")))
                continue
            if info.file_size > per_member:
                res.children.append((name, _fail("unsupported", ST_TOO_LARGE, f"Member is larger than {per_member:,} bytes.")))
                continue
            info_fmt = classify(name, None)
            if not info_fmt.supported:
                res.children.append((name, _fail(info_fmt.format, ST_UNSUPPORTED, info_fmt.reason or "Unsupported type.")))
                continue
            if info_fmt.format == "zip" and depth >= 1:
                res.children.append((name, _fail("zip", ST_UNSUPPORTED, "Nested archives are not opened.")))
                continue
            # Read with a hard cap: a member that lies about its size still cannot exceed the limit.
            try:
                with zf.open(info) as fh:
                    blob = fh.read(per_member + 1)
            except Exception as exc:
                res.children.append((name, _fail(info_fmt.format, ST_ERROR, f"Could not be extracted ({type(exc).__name__}).")))
                continue
            if len(blob) > per_member:
                res.children.append((name, _fail(info_fmt.format, ST_TOO_LARGE, "Member exceeded its size limit while extracting.")))
                continue
            read_total += len(blob)
            if read_total > max_total:
                res.children.append((name, _fail(info_fmt.format, ST_TOO_LARGE, "Archive total size limit reached; stopped.")))
                res.status = ST_PARTIAL
                break
            child = extract_bytes(blob, info_fmt.format, deadline=deadline, ocr=ocr, zip_depth=depth + 1)
            res.children.append((name, child))
    for name, child in res.children:
        for seg in child.segments:
            res.segments.append(Segment(seg.text, {"archive_member": name, **seg.ref}, seg.kind))
        for rec in child.records:
            res.records.append({"data": rec["data"], "ref": {"archive_member": name, **rec["ref"]}})
        for img in child.images:
            res.images.append({**img, "ref": {"archive_member": name, **(img.get("ref") or {})}})
    usable = [c for _, c in res.children if c.usable]
    res.meta.update({"members": len(res.children), "usable_members": len(usable)})
    if not res.children:
        res.status, res.error = ST_EMPTY, "The archive is empty."
    elif not usable:
        res.status, res.error = ST_EMPTY if res.status == ST_OK else res.status, "No supported, readable files inside the archive."
    elif len(usable) < len(res.children) and res.status == ST_OK:
        res.status = ST_PARTIAL
        res.warnings.append(f"{len(res.children) - len(usable)} archive member(s) could not be used")
    return res


# --------------------------------------------------------------------------------------------------------------------
# Dispatcher
# --------------------------------------------------------------------------------------------------------------------
def extract_bytes(data: bytes, fmt: str, *, deadline: Optional[Deadline] = None,
                  ocr: Optional[Callable[[bytes, str], str]] = None, zip_depth: int = 0) -> ExtractionResult:
    """Verify the signature, then run the extractor for the format. Never raises except ExtractionTimeout."""
    deadline = deadline or Deadline(None)
    if len(data) > max_file_bytes():
        return _fail(fmt, ST_TOO_LARGE, f"The file is larger than the {max_file_bytes():,}-byte limit.")
    use_fmt, problem = verify_signature(fmt, data)
    if problem:
        return _fail(fmt, problem[0], problem[1])
    label_fmt = fmt if fmt.startswith("google_") else use_fmt
    try:
        if use_fmt == "pdf":
            res = extract_pdf(data, deadline, ocr)
        elif use_fmt in ("docx", "google_doc"):
            res = extract_docx(data, deadline, fmt=label_fmt)
        elif use_fmt in ("xlsx", "google_sheet"):
            res = extract_xlsx(data, deadline, fmt=label_fmt)
        elif use_fmt == "xls":
            res = extract_xls(data, deadline)
        elif use_fmt == "csv":
            res = extract_csv(data, deadline)
        elif use_fmt in ("pptx", "google_slides"):
            res = extract_pptx(data, deadline, fmt=label_fmt)
        elif use_fmt in ("txt", "md"):
            res = extract_text(data, deadline, fmt=use_fmt)
        elif use_fmt == "rtf":
            res = extract_rtf(data, deadline)
        elif use_fmt == "html":
            res = extract_html(data, deadline)
        elif use_fmt == "json":
            res = extract_json(data, deadline)
        elif use_fmt == "xml":
            res = extract_xml(data, deadline)
        elif use_fmt == "image":
            res = extract_image(data, deadline, ocr)
        elif use_fmt == "zip":
            res = extract_zip(data, deadline, ocr, depth=zip_depth)
        else:
            res = _fail(use_fmt, ST_UNSUPPORTED, f"No extractor for format '{use_fmt}'.")
    except ExtractionTimeout:
        raise
    except MemoryError:
        res = _fail(use_fmt, ST_TOO_LARGE, "The file needed more memory than allowed to parse.")
    except Exception as exc:
        logger.warning("extractor %s failed: %s", use_fmt, type(exc).__name__)
        res = _fail(use_fmt, ST_ERROR, f"Parsing failed ({type(exc).__name__}).")
    if use_fmt != fmt and not fmt.startswith("google_"):
        res.warnings.append(f"declared as {fmt} but the content is {use_fmt}; parsed as {use_fmt}")
    return res
